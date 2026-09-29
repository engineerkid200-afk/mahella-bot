"""خواندن هوشمند اکسل/CSV با هر شکل جدول + خروجی اکسل از محصولات."""
import csv
import io
import json
import re
import statistics

from util import is_number, norm, to_int

FIELDS = ["code", "name", "brand", "category", "description", "price", "stock", "photo"]
FA = {"code": "کد", "name": "نام", "brand": "برند", "category": "دسته", "description": "توضیحات",
      "price": "قیمت", "stock": "موجودی", "photo": "لینک عکس"}

ALIASES = {
    "code": ["کد", "کد کالا", "کد محصول", "کدکالا", "کد کالا", "شناسه", "بارکد", "سریال", "ردیف کالا", "code", "sku", "id", "barcode"],
    "name": ["نام", "نام کالا", "نام محصول", "محصول", "کالا", "عنوان", "شرح", "شرح کالا", "name", "title", "product", "item"],
    "brand": ["برند", "نام برند", "مارک", "شرکت", "سازنده", "brand", "company", "maker"],
    "category": ["دسته", "دسته بندی", "گروه", "گروه کالا", "گروه بندی", "کتگوری", "category", "group"],
    "description": ["توضیح", "توضیحات", "شرح محصول", "مشخصات", "ویژگی", "ویژگی ها", "ویژگیها", "description", "desc", "details"],
    "price": ["قیمت", "قیمت فروش", "فی", "مبلغ", "قیمت واحد", "قیمت مصرف کننده", "قیمت تومان", "قیمت ریال",
              "بهای فروش", "price", "cost", "amount"],
    "stock": ["موجودی", "تعداد", "موجودی انبار", "مقدار", "تعداد موجود", "موجودی کالا", "stock", "qty", "quantity", "count", "inventory"],
    "photo": ["عکس", "تصویر", "لینک عکس", "لینک تصویر", "آدرس عکس", "photo", "image", "img", "picture", "url"],
}
CONTAINS_ORDER = ["code", "price", "stock", "brand", "category", "photo", "description", "name"]
IGNORED_HEADERS = {_k for _k in ("ردیف", "row", "#", "شماره", "no", "ش", "جمع", "total")}
GENERIC_SHEETS = re.compile(r"^(sheet|page|برگه|صفحه|کاربرگ)\s*\d*$", re.I)


def _key(s):
    return re.sub(r"[\s()\[\]\-_:/\\.*]+", "", norm(s))


AK = {f: [_key(a) for a in v] for f, v in ALIASES.items()}


def match_header(h):
    k = _key(h)
    if not k:
        return None
    for f in FIELDS:
        if k in AK[f]:
            return f
    for f in CONTAINS_ORDER:
        for a in AK[f]:
            if len(a) >= 3 and a in k:
                return f
    return None


def _empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def find_header(rows, limit=40):
    best_score, best_i = 0, None
    for i, r in enumerate(rows[:limit]):
        fields = {match_header(c) for c in r if isinstance(c, str) and c.strip()} - {None}
        if len(fields) > best_score:
            best_score, best_i = len(fields), i
    if best_score >= 2:
        return best_i
    if best_score == 1 and best_i is not None:
        only = {match_header(c) for c in rows[best_i] if isinstance(c, str)} - {None}
        if only == {"name"}:
            return best_i
    return None


def map_columns(header):
    mapping, used = {}, set()
    for i, h in enumerate(header):
        if _empty(h):
            continue
        k = _key(h)
        for f in FIELDS:
            if f not in mapping and k in AK[f]:
                mapping[f] = i
                used.add(i)
                break
    for i, h in enumerate(header):
        if i in used or _empty(h):
            continue
        f = match_header(h)
        if f and f not in mapping:
            mapping[f] = i
            used.add(i)
    extras = [i for i, h in enumerate(header)
              if i not in used and not _empty(h) and _key(h) not in IGNORED_HEADERS]
    return mapping, extras


def infer_columns(rows):
    """اگر ردیف عنوان نبود، از روی محتوای ستون‌ها حدس می‌زند."""
    rows = rows[:200]
    n = max(len(r) for r in rows)
    info = {}
    for i in range(n):
        vals = [r[i] for r in rows if i < len(r) and not _empty(r[i])]
        if len(vals) < max(1, int(len(rows) * 0.3)):
            continue
        nums = [v for v in vals if is_number(v)]
        isnum = len(nums) >= 0.8 * len(vals)
        strs = [str(v).strip() for v in vals]
        info[i] = {
            "isnum": isnum,
            "med": statistics.median([to_int(v, 0) for v in nums]) if isnum else 0,
            "distinct": len(set(strs)) / len(strs),
            "avg": sum(len(s) for s in strs) / len(strs),
        }
    m = {}
    nums = {i: d for i, d in info.items() if d["isnum"]}
    texts = {i: d for i, d in info.items() if not d["isnum"]}
    for i, d in sorted(nums.items()):
        if d["med"] > 1e8 and d["distinct"] > 0.95 and "code" not in m:
            m["code"] = i
    rest = {i: d for i, d in nums.items() if i != m.get("code")}
    if rest:
        price_i = max(rest, key=lambda i: rest[i]["med"])
        m["price"] = price_i
        others = [i for i in rest if i != price_i and rest[i]["med"] <= rest[price_i]["med"]]
        if others:
            m["stock"] = min(others)
    cand = {i: d for i, d in texts.items() if d["distinct"] > 0.6}
    if cand:
        m["name"] = max(cand, key=lambda i: cand[i]["avg"])
    left = {i: d for i, d in texts.items() if i != m.get("name")}
    if "code" not in m:
        for i, d in sorted(left.items()):
            if d["distinct"] > 0.95 and d["avg"] <= 14 and i < m.get("name", 99):
                m["code"] = i
                left.pop(i)
                break
    cat = [i for i, d in left.items() if d["distinct"] <= 0.4]
    if cat:
        m["category"] = min(cat, key=lambda i: left[i]["distinct"])
        left.pop(m["category"])
    if left:
        long_ = max(left, key=lambda i: left[i]["avg"])
        if left[long_]["avg"] > 40:
            m["description"] = long_
            left.pop(long_)
    if left:
        m["brand"] = min(left, key=lambda i: left[i]["distinct"])
    return m


def _code(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "".join(norm(v).split()).upper()


def _fmt(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_sheets(data, ext):
    if ext in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
                for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
    if ext == "xls":
        try:
            import xlrd
        except ImportError:
            raise ValueError("فایل xls قدیمی است؛ لطفاً در اکسل با Save As به xlsx تبدیلش کنید.")
        wb = xlrd.open_workbook(file_contents=data)
        return [(s.name, [s.row_values(i) for i in range(s.nrows)]) for s in wb.sheets() if getattr(s, "visibility", 0) == 0]
    text = None
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        text = data.decode("utf-16")
    else:
        for enc in ("utf-8-sig", "cp1256"):
            try:
                text = data.decode(enc)
                break
            except UnicodeError:
                pass
    if text is None:
        text = data.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [("csv", [r for r in csv.reader(io.StringIO(text), dialect)])]


def parse_file(data, filename):
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else "csv"
    sheets = read_sheets(data, ext)
    products, skipped, notes, mapping_lines = [], [], [], []
    extras_seen = set()
    stock_missing = False
    for title, rows in sheets:
        rows = [list(r) for r in rows]
        hi = find_header(rows)
        guessed = hi is None
        if not guessed:
            header = rows[hi]
            mapping, extras = map_columns(header)
            data_rows, start_no = rows[hi + 1:], hi + 2
        else:
            data_rows = [r for r in rows if any(not _empty(c) for c in r)]
            if len(data_rows) < 2:
                continue
            header, extras, start_no = None, [], 1
            mapping = infer_columns(data_rows)
        if "name" not in mapping or "price" not in mapping:
            cols = [str(h).strip() for h in (header or []) if not _empty(h)]
            missing = "نام" if "name" not in mapping else "قیمت"
            notes.append(f"برگه «{title}»: ستون {missing} پیدا نشد."
                         + (f" ستون‌ها: {' | '.join(cols)}" if cols else " (عنوان ستونی ندارد)"))
            continue
        if guessed:
            mapping_lines.append(f"«{title}» (بدون عنوان، حدسی): " + "، ".join(f"{FA[f]}=ستون {i + 1}" for f, i in mapping.items()))
        else:
            mapping_lines.append(f"«{title}»: " + "، ".join(f"{FA[f]}←{str(header[i]).strip()}" for f, i in mapping.items()))
        if "stock" not in mapping:
            stock_missing = True
        rial = (not guessed) and any(w in norm(header[mapping["price"]]) for w in ("ریال", "rial", "irr"))
        if rial:
            notes.append(f"برگه «{title}»: قیمت به ریال بود، به تومان تبدیل شد.")
        default_cat = None if "category" in mapping or GENERIC_SHEETS.match(title.strip()) or title == "csv" else title.strip()
        cur_cat = None

        def cell(r, f):
            i = mapping.get(f)
            return r[i] if i is not None and i < len(r) else None

        for off, r in enumerate(data_rows):
            no = start_no + off
            filled = [c for c in r if not _empty(c)]
            if not filled:
                continue
            if len(filled) == 1 and isinstance(filled[0], str) and not is_number(filled[0]):
                cur_cat = filled[0].strip()
                continue
            name = _fmt(cell(r, "name")) if not _empty(cell(r, "name")) else ""
            if not name:
                skipped.append((no, "بدون نام"))
                continue
            price = to_int(cell(r, "price"))
            if price is None:
                skipped.append((no, f"قیمت خالی یا نامعتبر ({name[:20]})"))
                continue
            if rial:
                price = int(round(price / 10))
            if price <= 0:
                skipped.append((no, f"قیمت صفر ({name[:20]})"))
                continue
            raw_code = cell(r, "code")
            sv = cell(r, "stock")
            stock = None if _empty(sv) else max(0, to_int(sv, 0))
            pv = cell(r, "photo")
            photo = str(pv).strip() if isinstance(pv, str) and pv.strip().lower().startswith("http") else None
            cat = cell(r, "category")
            cat = _fmt(cat) if not _empty(cat) else (cur_cat or default_cat)
            extra = {}
            for i in extras:
                v = r[i] if i < len(r) else None
                if not _empty(v) and len(extra) < 12:
                    extra[str(header[i]).strip()] = _fmt(v)[:200]
                    extras_seen.add(str(header[i]).strip())
            products.append({
                "code": None if _empty(raw_code) else _code(raw_code) or None,
                "name": name,
                "brand": None if _empty(cell(r, "brand")) else _fmt(cell(r, "brand")),
                "category": cat or None,
                "description": None if _empty(cell(r, "description")) else _fmt(cell(r, "description")),
                "price": price, "stock": stock, "photo": photo, "extra": extra,
            })
    return {"products": products, "skipped": skipped, "notes": notes, "mapping": mapping_lines,
            "extras": sorted(extras_seen), "stock_missing": stock_missing}


def export_xlsx(products):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "محصولات"
    ws.sheet_view.rightToLeft = True
    extra_keys = []
    for p in products:
        for k in json.loads(p.get("extra") or "{}"):
            if k not in extra_keys:
                extra_keys.append(k)
    ws.append(["کد", "نام", "برند", "دسته", "توضیحات", "قیمت", "موجودی", "لینک عکس"] + extra_keys)
    for c in ws[1]:
        c.font = Font(bold=True)
    for p in products:
        ex = json.loads(p.get("extra") or "{}")
        photo = p["photo"] if str(p.get("photo") or "").startswith("http") else ""
        ws.append([p["code"], p["name"], p["brand"], p["category"], p["description"], p["price"], p["stock"], photo]
                  + [ex.get(k, "") for k in extra_keys])
    for col, w in zip("ABCDEFGH", (14, 32, 16, 18, 40, 12, 10, 24)):
        ws.column_dimensions[col].width = w
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
